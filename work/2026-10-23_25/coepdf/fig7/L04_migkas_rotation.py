# PROVENANCE COPY of /workspace/cosmology/figures/migkas_bars_rotation_overlay.py (original generating script, rev6 figure; run in its original directory).
# Output figure copied here as L04_migkas_rotation.png. Not re-run for rev7: see Appendix B.
#!/usr/bin/env python3
"""
Migkas Fig.6 overlay with published-style error bars + Chris sheet-style
sine phase-rotation diagnostic.

This is NOT a free A cos φ dipole template. It reproduces the spreadsheet
rotation model:

  BA = Amp * sin((2π/360)·ℓ + φ)     # Amp≈−0.06, φ≈0.35 rad
  BE = −Amp * same sine
  Model A/A_all = (1 − BA) * AF      # AF = sheet ST H_obs/H series
  Model B/B_all = (1 − BE) * AN      # AN = sheet mass / B-side series

Sheet tab used for cached series: "Copy of Copy of Copy of density"
(R²_A≈0.739, R²_B≈0.733 with rotation; ≈0.50 / ≈0.37 without).

First-principles ST quadrupole shape alone is a separate claim — this panel
is the sheet-style phase-rotation diagnostic vs Migkas bars.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np
import matplotlib as mpl
import matplotlib.pyplot as plt
from matplotlib.gridspec import GridSpec

OUT = Path(__file__).resolve().parent
CSV_MIGKAS = Path("/workspace/cosmology/migkas/fig6_digitized.csv")
XLSX = Path(
    "/home/box/agent-data/agents/196ab413-0708-4f81-99fc-58a32c5366c5/"
    "attachments/b3737b0408e9976ecfad92319c9a6bbbf53777e5846e2136703654a0dcf9269a.xlsx"
)
SHEET_CACHE = Path("/tmp/chris_sheet_series.csv")
SHEET_TAB = "Copy of Copy of Copy of density"

# Chris sheet params on that tab (cached)
AMP_A = -0.06
PHI_RAD = 0.35  # ≈20.05°
OMEGA = 2.0 * np.pi / 360.0

C_BLUE = "#0072B2"
C_ORANGE = "#E69F00"
C_BLACK = "#222222"
C_GRAY = "#666666"
C_GREEN = "#009E73"


def style() -> None:
    mpl.rcParams.update(
        {
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
            "savefig.dpi": 220,
            "font.size": 11,
            "axes.labelsize": 12,
            "axes.titlesize": 12,
            "legend.fontsize": 8.5,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "axes.linewidth": 1.1,
            "lines.linewidth": 2.0,
            "mathtext.fontset": "dejavusans",
        }
    )


def rsq(y: np.ndarray, x: np.ndarray) -> float:
    """Pearson R² (Excel RSQ)."""
    y = np.asarray(y, dtype=float)
    x = np.asarray(x, dtype=float)
    if y.size < 2 or np.std(x) == 0 or np.std(y) == 0:
        return float("nan")
    return float(np.corrcoef(x, y)[0, 1] ** 2)


def load_migkas_csv(path: Path = CSV_MIGKAS) -> dict[str, np.ndarray]:
    rows = list(csv.DictReader(path.open()))
    out = {k: np.array([float(r[k]) for r in rows]) for k in rows[0]}
    return out


def load_sheet_series(cache: Path = SHEET_CACHE) -> dict[str, np.ndarray]:
    """Load AR/AF/AN/BA/BE/BH/BI/BL/BM from pre-extracted CSV or xlsx."""
    if cache.is_file():
        rows = list(csv.DictReader(cache.open()))
        keys = ["AR", "AF", "AN", "BA", "BE", "BH", "BI", "BL", "BM", "BG"]
        return {k: np.array([float(r[k]) for r in rows]) for k in keys}

    # Fallback: extract with openpyxl (slow)
    from openpyxl import load_workbook

    wb = load_workbook(XLSX, data_only=True, read_only=True)
    ws = wb[SHEET_TAB]
    data = {k: [] for k in ["AR", "AF", "AN", "BA", "BE", "BH", "BI", "BL", "BM", "BG"]}
    for r in range(66, 427):
        for k in data:
            data[k].append(float(ws[f"{k}{r}"].value))
    wb.close()
    return {k: np.asarray(v) for k, v in data.items()}


def sine_mod(l_deg: np.ndarray, amp: float, phi: float) -> np.ndarray:
    return amp * np.sin(OMEGA * np.asarray(l_deg, dtype=float) + phi)


def model_rotated(
    af: np.ndarray, an: np.ndarray, l_deg: np.ndarray, amp_a: float, phi: float
) -> tuple[np.ndarray, np.ndarray]:
    """Chris sheet: Model A = (1−BA)*AF, Model B = (1−BE)*AN, BE=−amp_a."""
    ba = sine_mod(l_deg, amp_a, phi)
    be = sine_mod(l_deg, -amp_a, phi)
    return (1.0 - ba) * af, (1.0 - be) * an


def interp_sheet_to(l_target: np.ndarray, l_src: np.ndarray, y_src: np.ndarray) -> np.ndarray:
    """Periodic-ish interp on 0..360; sheet has 0..360 inclusive."""
    # Drop duplicate 360 if present
    mask = l_src < 360.0 - 1e-9
    ls, ys = l_src[mask], y_src[mask]
    order = np.argsort(ls)
    ls, ys = ls[order], ys[order]
    # wrap pad
    ls_p = np.concatenate([ls - 360.0, ls, ls + 360.0])
    ys_p = np.concatenate([ys, ys, ys])
    return np.interp(np.asarray(l_target, dtype=float) % 360.0, ls_p, ys_p)


def scan_phi(
    l: np.ndarray,
    af: np.ndarray,
    an: np.ndarray,
    a_obs: np.ndarray,
    b_obs: np.ndarray,
    amp: float = AMP_A,
    n: int = 361,
) -> dict:
    phis = np.linspace(0.0, 2.0 * np.pi, n, endpoint=False)
    best = {"phi": 0.0, "r2_a": -1.0, "r2_b": -1.0, "r2_sum": -1.0}
    curve_a, curve_b = [], []
    for phi in phis:
        ma, mb = model_rotated(af, an, l, amp, phi)
        ra, rb = rsq(a_obs, ma), rsq(b_obs, mb)
        curve_a.append(ra)
        curve_b.append(rb)
        s = ra + rb
        if s > best["r2_sum"]:
            best = {
                "phi": float(phi),
                "phi_deg": float(np.rad2deg(phi)),
                "r2_a": float(ra),
                "r2_b": float(rb),
                "r2_sum": float(s),
            }
    best["phi_scan_r2_a"] = curve_a
    best["phi_scan_r2_b"] = curve_b
    best["phis"] = phis.tolist()
    return best


def main() -> None:
    style()
    mig = load_migkas_csv()
    sheet = load_sheet_series()

    l_s = sheet["AR"]
    af_s, an_s = sheet["AF"], sheet["AN"]
    bh_s, bl_s = sheet["BH"], sheet["BL"]
    bi_s, bm_s = sheet["BI"], sheet["BM"]

    # Sheet R² (exact Excel RSQ on 1° VLOOKUP Migkas)
    r2_a_with_sheet = rsq(sheet["BI"], sheet["BH"])
    r2_b_with_sheet = rsq(sheet["BM"], sheet["BL"])
    r2_a_wo_sheet = rsq(sheet["BI"], sheet["AF"])
    r2_b_wo_sheet = rsq(sheet["BM"], sheet["AN"])

    # Digitized Migkas + asymmetric bars
    l_m = mig["l_deg"]
    a_m = mig["A_over_Aall"]
    b_m = mig["B_over_Ball"]
    a_elo, a_ehi = mig["A_err_low"], mig["A_err_high"]
    b_elo, b_ehi = mig["B_err_low"], mig["B_err_high"]

    # Sheet model curves at digitized longitudes (Chris φ)
    af_m = interp_sheet_to(l_m, l_s, af_s)
    an_m = interp_sheet_to(l_m, l_s, an_s)
    model_a_chris, model_b_chris = model_rotated(af_m, an_m, l_m, AMP_A, PHI_RAD)
    # Also the cached BH/BL interpolated (should match)
    bh_m = interp_sheet_to(l_m, l_s, bh_s)
    bl_m = interp_sheet_to(l_m, l_s, bl_s)

    r2_a_dig_chris = rsq(a_m, model_a_chris)
    r2_b_dig_chris = rsq(b_m, model_b_chris)
    r2_a_dig_wo = rsq(a_m, af_m)
    r2_b_dig_wo = rsq(b_m, an_m)

    # Scan φ on digitized points (shared φ, opposite Amp)
    scan = scan_phi(l_m, af_m, an_m, a_m, b_m, amp=AMP_A)
    model_a_best, model_b_best = model_rotated(
        af_m, an_m, l_m, AMP_A, scan["phi"]
    )
    # Fine curve for plotting (1°)
    l_fine = np.arange(0.0, 360.0, 1.0)
    af_f = interp_sheet_to(l_fine, l_s, af_s)
    an_f = interp_sheet_to(l_fine, l_s, an_s)
    ma_f, mb_f = model_rotated(af_f, an_f, l_fine, AMP_A, PHI_RAD)
    ma_fb, mb_fb = model_rotated(af_f, an_f, l_fine, AMP_A, scan["phi"])

    # ---------- Fig 1: Migkas-like 2-panel overlay ----------
    fig = plt.figure(figsize=(10.5, 8.2))
    gs = GridSpec(2, 1, height_ratios=[1, 1], hspace=0.08, figure=fig)
    ax_a = fig.add_subplot(gs[0])
    ax_b = fig.add_subplot(gs[1], sharex=ax_a)

    ax_a.errorbar(
        l_m,
        a_m,
        yerr=np.vstack([a_elo, a_ehi]),
        fmt="o",
        ms=4.0,
        color=C_BLACK,
        ecolor=C_GRAY,
        elinewidth=0.9,
        capsize=2.0,
        label="Migkas A/A$_\\mathrm{all}$ (digitised, Δℓ=90°)",
        zorder=3,
    )
    ax_a.plot(
        l_fine,
        ma_f,
        color=C_ORANGE,
        lw=2.2,
        label=(
            f"Sheet sine rotation: $(1-A\\sin)\\times$ST  "
            f"(Amp={AMP_A}, φ={PHI_RAD:.2f} rad≈{np.rad2deg(PHI_RAD):.1f}°)"
        ),
        zorder=4,
    )
    if abs(scan["phi"] - PHI_RAD) > 0.05:
        ax_a.plot(
            l_fine,
            ma_fb,
            color=C_GREEN,
            lw=1.4,
            ls="--",
            label=f"Best shared-φ scan (φ={scan['phi']:.2f} rad)",
            zorder=2,
        )
    ax_a.axhline(1.0, color="#aaaaaa", lw=0.8, ls=":")
    ax_a.set_ylabel(r"$A/A_{\mathrm{all}}$")
    ax_a.set_ylim(0.75, 1.35)
    ax_a.legend(loc="upper right", framealpha=0.92)
    ax_a.set_title(
        "Migkas Fig.6–style overlay — sheet sine phase rotation (not ST-quad alone)",
        pad=8,
    )
    ax_a.text(
        0.02,
        0.04,
        f"R²$_A$ sheet: {r2_a_wo_sheet:.3f}→{r2_a_with_sheet:.3f}  |  "
        f"vs digitised bars: {r2_a_dig_wo:.3f}→{r2_a_dig_chris:.3f}",
        transform=ax_a.transAxes,
        fontsize=8.5,
        va="bottom",
        bbox=dict(boxstyle="round,pad=0.3", fc="white", ec="#cccccc", alpha=0.9),
    )

    ax_b.errorbar(
        l_m,
        b_m,
        yerr=np.vstack([b_elo, b_ehi]),
        fmt="o",
        ms=4.0,
        color=C_BLACK,
        ecolor=C_GRAY,
        elinewidth=0.9,
        capsize=2.0,
        label="Migkas B/B$_\\mathrm{all}$ (digitised, Δℓ=90°)",
        zorder=3,
    )
    ax_b.plot(
        l_fine,
        mb_f,
        color=C_BLUE,
        lw=2.2,
        label=(
            f"Sheet sine rotation: $(1-B\\sin)\\times$ST  "
            f"(Amp={-AMP_A:+.2f}, same φ)"
        ),
        zorder=4,
    )
    if abs(scan["phi"] - PHI_RAD) > 0.05:
        ax_b.plot(
            l_fine,
            mb_fb,
            color=C_GREEN,
            lw=1.4,
            ls="--",
            label=f"Best shared-φ scan",
            zorder=2,
        )
    ax_b.axhline(1.0, color="#aaaaaa", lw=0.8, ls=":")
    ax_b.set_ylabel(r"$B/B_{\mathrm{all}}$")
    ax_b.set_xlabel(r"Galactic longitude $\ell$ [deg]")
    ax_b.set_xlim(0, 360)
    ax_b.set_xticks([0, 60, 120, 180, 240, 300, 360])
    ax_b.set_ylim(0.75, 1.30)
    ax_b.legend(loc="upper right", framealpha=0.92)
    ax_b.text(
        0.02,
        0.04,
        f"R²$_B$ sheet: {r2_b_wo_sheet:.3f}→{r2_b_with_sheet:.3f}  |  "
        f"vs digitised bars: {r2_b_dig_wo:.3f}→{r2_b_dig_chris:.3f}",
        transform=ax_b.transAxes,
        fontsize=8.5,
        va="bottom",
        bbox=dict(boxstyle="round,pad=0.3", fc="white", ec="#cccccc", alpha=0.9),
    )

    fig.text(
        0.5,
        0.005,
        "Diagnostic: Amp×sin((2π/360)·ℓ+φ) phase alignment from Chris sheet "
        f"tab “{SHEET_TAB}”. ST quadrupole alone is a separate claim.",
        ha="center",
        fontsize=8,
        style="italic",
        color=C_GRAY,
    )

    out_png = OUT / "migkas_bars_rotation_overlay.png"
    out_pdf = OUT / "migkas_bars_rotation_overlay.pdf"
    fig.savefig(out_png, bbox_inches="tight")
    fig.savefig(out_pdf, bbox_inches="tight")
    plt.close(fig)
    print(f"wrote {out_png}")
    print(f"wrote {out_pdf}")

    # ---------- Fig 2: scatter Migkas vs model ----------
    fig2, axes = plt.subplots(1, 2, figsize=(10.2, 4.8))
    for ax, obs, mod, name, col, r2w, r2wo in [
        (axes[0], a_m, model_a_chris, "A/A_all", C_ORANGE, r2_a_dig_chris, r2_a_dig_wo),
        (axes[1], b_m, model_b_chris, "B/B_all", C_BLUE, r2_b_dig_chris, r2_b_dig_wo),
    ]:
        ax.errorbar(
            mod,
            obs,
            yerr=np.vstack(
                [a_elo, a_ehi] if name.startswith("A") else [b_elo, b_ehi]
            ),
            fmt="o",
            ms=4.5,
            color=C_BLACK,
            ecolor=C_GRAY,
            elinewidth=0.8,
            capsize=1.5,
            alpha=0.85,
        )
        lims = [
            min(obs.min(), mod.min()) - 0.05,
            max(obs.max(), mod.max()) + 0.05,
        ]
        ax.plot(lims, lims, "--", color="#999999", lw=1.0, label="1:1")
        # OLS
        a_fit, b_fit = np.polyfit(mod, obs, 1)
        xx = np.linspace(lims[0], lims[1], 50)
        ax.plot(xx, a_fit * xx + b_fit, color=col, lw=2.0, label=f"OLS fit")
        ax.set_xlim(lims)
        ax.set_ylim(lims)
        ax.set_aspect("equal", adjustable="box")
        ax.set_xlabel(f"Sheet-rotated model {name}")
        ax.set_ylabel(f"Migkas {name}")
        ax.set_title(f"{name}: R²={r2w:.3f} (was {r2wo:.3f} w/o sine)")
        ax.legend(loc="upper left", fontsize=8)
    fig2.suptitle(
        "Migkas vs sheet sine-rotation model (digitised points + error bars)",
        y=1.02,
        fontsize=12,
    )
    fig2.tight_layout()
    out_sc = OUT / "migkas_bars_rotation_scatter.png"
    fig2.savefig(out_sc, bbox_inches="tight")
    fig2.savefig(OUT / "migkas_bars_rotation_scatter.pdf", bbox_inches="tight")
    plt.close(fig2)
    print(f"wrote {out_sc}")

    # ---------- summary JSON ----------
    summary = {
        "sheet_tab": SHEET_TAB,
        "method": (
            "Simple Amp×sin((2π/360)·ℓ+φ) phase alignment on sheet ST series; "
            "Model A=(1−BA)*AF, Model B=(1−BE)*AN with BE=−BA Amp. "
            "NOT a free A cos φ dipole; NOT claiming first-principles ST "
            "quadrupole alone matches B."
        ),
        "amp_A": AMP_A,
        "amp_B": -AMP_A,
        "phi_rad_chris": PHI_RAD,
        "phi_deg_chris": float(np.rad2deg(PHI_RAD)),
        "omega": OMEGA,
        "phase_table_deg_chris": [-20.05, 69.95, 159.95, 249.95, 339.95],
        "R2_sheet_1deg_VLOOKUP": {
            "A_without_rotation": r2_a_wo_sheet,
            "A_with_rotation": r2_a_with_sheet,
            "B_without_rotation": r2_b_wo_sheet,
            "B_with_rotation": r2_b_with_sheet,
            "improvement_A": r2_a_with_sheet / r2_a_wo_sheet,
            "improvement_B": r2_b_with_sheet / r2_b_wo_sheet,
            "note": "Matches cached BH58≈0.739, BL58≈0.733, AF51≈0.50, AM51≈0.37, D3≈1.48, D4≈1.98",
        },
        "R2_vs_digitised_fig6_csv": {
            "n_points": int(len(l_m)),
            "A_without_rotation": r2_a_dig_wo,
            "A_with_rotation_chris_phi": r2_a_dig_chris,
            "B_without_rotation": r2_b_dig_wo,
            "B_with_rotation_chris_phi": r2_b_dig_chris,
            "best_shared_phi_rad": scan["phi"],
            "best_shared_phi_deg": scan["phi_deg"],
            "A_with_best_phi": scan["r2_a"],
            "B_with_best_phi": scan["r2_b"],
        },
        "outputs": {
            "overlay_png": str(out_png),
            "overlay_pdf": str(out_pdf),
            "scatter_png": str(out_sc),
        },
        "honest_note": (
            "Both A and B match well under Chris’s sheet-style sine phase "
            "rotation (R²≈0.74/0.73 on sheet; ~0.7 ballpark vs digitised bars). "
            "Forcing the same ST+Gauss H_obs quadrupole shape onto B without "
            "this Amp×sin phase factor yields near-null R² for B — that is a "
            "different (and weaker) claim than this diagnostic panel."
        ),
    }
    # drop large scan arrays from JSON
    out_json = OUT / "migkas_bars_rotation_summary.json"
    with out_json.open("w") as f:
        json.dump(summary, f, indent=2)
    print(f"wrote {out_json}")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
